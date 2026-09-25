#!/usr/bin/env python3
"""Product部門方針 — 編集可能な PowerPoint を生成する（画像貼り付けではなくテキスト・表）。

Google スライドで編集する場合:
  ⚠️ ドライブに .pptx を上げてダブルクリックすると「ファイルを開けませんでした」になる。
     PPTX は Google スライド形式ではないため、次のいずれかで開く:

  【手動・インポート】
    product-dept-slides.pptx を使う（日本語ファイル名より通りやすい）
    1. https://slides.google.com で「空白」を新規作成
    2. ファイル → スライドをインポート → product-dept-slides.pptx
    ※「アップロード」タブで失敗する場合:
       drive.google.com に先にアップロード → インポートの「マイドライブ」タブから選択

  【自動・ネイティブ形式】
    gcloud auth application-default login \\
      --scopes=https://www.googleapis.com/auth/presentations,https://www.googleapis.com/auth/drive
    python3 create-google-slides-native.py
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from slide_content import SLIDE1 as SLIDE1_RAW, SLIDE2 as SLIDE2_RAW

DIR = Path(__file__).resolve().parent
OUT = DIR / "Product部門方針.pptx"


def _hex_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _with_pptx_colors(data: dict) -> dict:
    out = dict(data)
    if "teams" in out:
        out["teams"] = [{**t, "color": _hex_rgb(t["color"])} for t in out["teams"]]
    if "cross" in out:
        out["cross"] = {**out["cross"], "color": _hex_rgb(out["cross"]["color"])}
    return out


SLIDE1 = _with_pptx_colors(SLIDE1_RAW)
SLIDE2 = _with_pptx_colors(SLIDE2_RAW)


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _text_box(slide, left, top, width, height, text, *, size=12, bold=False, color=None, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    if color:
        p.font.color.rgb = color
    p.alignment = align
    return box


def _fill_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def _set_cell(cell, text, *, size=9, bold=False, color=RGBColor(0x1A, 0x1A, 0x1A), align=PP_ALIGN.LEFT):
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(1)
    tf.margin_bottom = Pt(1)
    for p in tf.paragraphs:
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align


def _fill_cell(cell, color: RGBColor):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def _mission_table(slide, left, top, width, height, team: dict):
    rows_data = team["rows"]
    n_rows = 1 + len(rows_data)
    table = slide.shapes.add_table(n_rows, 3, left, top, width, height).table

    col_w = width / 3
    table.columns[0].width = int(col_w * 1.45)
    table.columns[1].width = int(col_w * 0.775)
    table.columns[2].width = int(col_w * 0.775)

    headers = ("ミッション", "現", "1年後")
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        _set_cell(cell, h, size=8, bold=True)
        _fill_cell(cell, RGBColor(0xF5, 0xF5, 0xF5))

    for r, (mission, now, future) in enumerate(rows_data, start=1):
        _set_cell(table.cell(r, 0), mission, size=8, bold=True)
        _fill_cell(table.cell(r, 0), RGBColor(0xFA, 0xFA, 0xFA))
        _set_cell(table.cell(r, 1), now, size=8, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(r, 2), future, size=8, align=PP_ALIGN.CENTER)

    return table


def _team_header(slide, left, top, width, team: dict):
    h = Inches(0.55)
    _fill_rect(slide, left, top, width, h, team["color"])
    _text_box(
        slide,
        left + Inches(0.08),
        top + Inches(0.04),
        width - Inches(0.16),
        Inches(0.22),
        f"{team['name']}　{team['tag']}",
        size=11,
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    _text_box(
        slide,
        left + Inches(0.08),
        top + Inches(0.26),
        width - Inches(0.16),
        Inches(0.26),
        team["outcome"],
        size=7,
        color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    return h


def build_slide1(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    data = SLIDE1
    margin = Inches(0.35)
    w = prs.slide_width - margin * 2

    _text_box(slide, margin, Inches(0.2), w, Inches(0.45), data["title"], size=24, bold=True)
    _text_box(slide, margin, Inches(0.62), w, Inches(0.35), data["vision"], size=14, bold=True, color=RGBColor(0x2F, 0x6F, 0xED))
    _text_box(slide, margin, Inches(0.95), w, Inches(0.3), data["vision_note"], size=10, color=RGBColor(0x55, 0x55, 0x55))

    col_w = (w - Inches(0.2)) / 3
    top = Inches(1.35)
    col_h = Inches(2.55)
    gap = Inches(0.1)

    for i, team in enumerate(data["teams"]):
        left = margin + (col_w + gap) * i
        _fill_rect(slide, left, top, col_w, col_h, RGBColor(0xFF, 0xFF, 0xFF), team["color"])
        hdr_h = _team_header(slide, left, top, col_w, team)
        body_top = top + hdr_h + Inches(0.08)
        bullets = "\n".join(f"• {b}" for b in team["bullets"])
        _text_box(slide, left + Inches(0.1), body_top, col_w - Inches(0.2), col_h - hdr_h - Inches(0.15), bullets, size=9)

    band_top = Inches(4.05)
    band_h = Inches(0.72)
    band_gap = Inches(0.08)

    for band in (data["cross_top"], data["cross_bottom"]):
        _fill_rect(slide, margin, band_top, w, band_h, RGBColor(0xF3, 0xF0, 0xF8), RGBColor(0x6B, 0x5B, 0x95))
        _text_box(slide, margin + Inches(0.12), band_top + Inches(0.06), w * 0.55, Inches(0.22), band["label"], size=10, bold=True, color=RGBColor(0x6B, 0x5B, 0x95))
        _text_box(slide, margin + Inches(0.12), band_top + Inches(0.28), w * 0.55, Inches(0.18), band["items"], size=8, color=RGBColor(0x44, 0x44, 0x44))
        _text_box(slide, margin + w * 0.58, band_top + Inches(0.12), w * 0.38, Inches(0.4), band["members"], size=11, bold=True, align=PP_ALIGN.RIGHT)
        band_top += band_h + band_gap

    _fill_rect(slide, margin, Inches(5.65), w, Inches(0.55), RGBColor(0x2F, 0x6F, 0xED))
    _text_box(slide, margin + Inches(0.12), Inches(5.75), w - Inches(0.24), Inches(0.4), data["goal_2027"], size=12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER)

    slide.notes_slide.notes_text_frame.text = "部門Vision・3チーム体制・横断役割・2027年ゴール（PowerPoint上で直接編集可）"


def build_slide2(prs: Presentation) -> None:
    slide = _blank_slide(prs)
    data = SLIDE2
    margin = Inches(0.3)
    w = prs.slide_width - margin * 2

    _text_box(slide, margin, Inches(0.15), w, Inches(0.4), data["title"], size=20, bold=True)
    _text_box(slide, margin, Inches(0.52), w, Inches(0.25), data["subtitle"], size=8, color=RGBColor(0x55, 0x55, 0x55))

    col_w = (w - Inches(0.16)) / 3
    top = Inches(0.82)
    hdr_h = Inches(0.55)
    tbl_h = Inches(2.35)
    gap = Inches(0.08)

    for i, team in enumerate(data["teams"]):
        left = margin + (col_w + gap) * i
        _fill_rect(slide, left, top, col_w, hdr_h + tbl_h, RGBColor(0xFF, 0xFF, 0xFF), team["color"])
        _team_header(slide, left, top, col_w, team)
        _mission_table(slide, left, top + hdr_h, col_w, tbl_h, team)

    cross = data["cross"]
    cross_top = top + hdr_h + tbl_h + Inches(0.12)
    cross_hdr_h = Inches(0.5)
    cross_tbl_h = Inches(1.55)

    _fill_rect(slide, margin, cross_top, w, cross_hdr_h + cross_tbl_h, RGBColor(0xFF, 0xFF, 0xFF), cross["color"])
    _fill_rect(slide, margin, cross_top, w, cross_hdr_h, cross["color"])
    _text_box(
        slide,
        margin + Inches(0.1),
        cross_top + Inches(0.04),
        w * 0.35,
        Inches(0.22),
        f"{cross['name']}　{cross['tag']}",
        size=11,
        bold=True,
        color=RGBColor(0xFF, 0xFF, 0xFF),
    )
    _text_box(
        slide,
        margin + Inches(0.1),
        cross_top + Inches(0.26),
        w - Inches(0.2),
        Inches(0.2),
        cross["outcome"],
        size=7,
        color=RGBColor(0xFF, 0xFF, 0xFF),
    )

    half_w = (w - Inches(0.06)) / 2
    tbl_top = cross_top + cross_hdr_h
    left_rows = [(m, n, f) for m, n, f in cross["rows_left"]]
    right_rows = [(m, n, f) for m, n, f in cross["rows_right"]]

    _mission_table(slide, margin, tbl_top, half_w, cross_tbl_h, {"rows": left_rows})
    _mission_table(slide, margin + half_w + Inches(0.06), tbl_top, half_w, cross_tbl_h, {"rows": right_rows})

    slide.notes_slide.notes_text_frame.text = "現時点（2026-08）と1年後の担当マッピング（PowerPoint上で直接編集可）"


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_slide1(prs)
    build_slide2(prs)

    prs.save(OUT)
    print(f"Created: {OUT}")

    # Google スライド向け互換版も生成
    subprocess.run([sys.executable, str(DIR / "build-google-import-pptx.py")], check=True)

    print("  → PowerPoint / Keynote で編集: Product部門方針.pptx")
    print("  → Google スライドでインポート: product-dept-slides.pptx（ASCII名・互換版）")
    print("  → インポート手順:")
    print("       1) slides.google.com で空白を作成")
    print("       2) ファイル → スライドをインポート")
    print("       3) product-dept-slides.pptx を選択")
    print("     ※「アップロード」で失敗する場合 → drive.google.com に先に上げ、")
    print("       インポート画面の「マイドライブ」タブから選ぶ")
    print("  → または: python3 create-google-slides-native.py（要 gcloud 認証）")
    print("  → 内容を変えるときは slide_content.py も更新してください。")


if __name__ == "__main__":
    main()
