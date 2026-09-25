#!/usr/bin/env python3
"""Google スライドの「スライドをインポート」向け互換 PPTX を生成する。

- product-dept-slides.pptx       … テキストのみ（インポート成功率優先）
- product-dept-slides-tables.pptx … 表あり（2枚目ミッションと担当）
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from slide_content import SLIDE1 as SLIDE1_RAW, SLIDE2 as SLIDE2_RAW

DIR = Path(__file__).resolve().parent
OUT_TEXT = DIR / "product-dept-slides.pptx"
OUT_TABLES = DIR / "product-dept-slides-tables.pptx"

SKIP_ZIP_ENTRIES = (
    "docProps/thumbnail.jpeg",
)


def _hex_rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _with_pptx_colors(data: dict) -> dict:
    out = dict(data)
    if "teams" in out and data["teams"] and "bullets" in data["teams"][0]:
        out["teams"] = [{**t, "color": _hex_rgb(t["color"])} for t in out["teams"]]
    elif "teams" in out:
        out["teams"] = [{**t, "color": _hex_rgb(t["color"])} for t in out["teams"]]
    if "cross" in out:
        out["cross"] = {**out["cross"], "color": _hex_rgb(out["cross"]["color"])}
    return out


SLIDE1 = _with_pptx_colors(SLIDE1_RAW)
SLIDE2 = _with_pptx_colors(SLIDE2_RAW)


def _text(slide, left, top, width, height, text, *, size=11, bold=False, color=None, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    if color:
        p.font.color.rgb = color
    p.alignment = align
    return box


def _rows_to_text(rows: list[tuple[str, str, str]]) -> str:
    lines = ["ミッション\t現\t1年後"]
    for mission, now, future in rows:
        lines.append(f"{mission}\t{now}\t{future}")
    return "\n".join(lines)


def build_slide1(slide) -> None:
    data = SLIDE1
    margin = Inches(0.3)
    w = Inches(10) - margin * 2

    _text(slide, margin, Inches(0.15), w, Inches(0.35), data["title"], size=20, bold=True)
    _text(slide, margin, Inches(0.48), w, Inches(0.3), data["vision"], size=12, bold=True, color=_hex_rgb("#2f6fed"))
    _text(slide, margin, Inches(0.75), w, Inches(0.22), data["vision_note"], size=8, color=RGBColor(0x55, 0x55, 0x55))

    col_w = (w - Inches(0.12)) / 3
    top = Inches(1.0)
    gap = Inches(0.06)

    for i, team in enumerate(data["teams"]):
        left = margin + (col_w + gap) * i
        header = f"{team['name']}（{team['tag']}）\n{team['outcome']}"
        body = "\n".join(f"・{b}" for b in team["bullets"])
        _text(slide, left, top, col_w, Inches(0.55), header, size=8, bold=True, color=team["color"])
        _text(slide, left, top + Inches(0.55), col_w, Inches(1.35), body, size=7)

    y = Inches(2.95)
    for band in (data["cross_top"], data["cross_bottom"]):
        _text(slide, margin, y, w * 0.65, Inches(0.38), f"{band['label']}\n{band['items']}", size=7)
        _text(slide, margin + w * 0.65, y, w * 0.35, Inches(0.38), band["members"], size=8, bold=True, align=PP_ALIGN.RIGHT)
        y += Inches(0.42)

    _text(slide, margin, Inches(4.75), w, Inches(0.35), data["goal_2027"], size=10, bold=True, color=_hex_rgb("#2f6fed"), align=PP_ALIGN.CENTER)


def build_slide2(slide) -> None:
    data = SLIDE2
    margin = Inches(0.25)
    w = Inches(10) - margin * 2

    _text(slide, margin, Inches(0.1), w, Inches(0.28), data["title"], size=16, bold=True)
    _text(slide, margin, Inches(0.36), w, Inches(0.18), data["subtitle"], size=6, color=RGBColor(0x55, 0x55, 0x55))

    col_w = (w - Inches(0.1)) / 3
    top = Inches(0.55)
    gap = Inches(0.05)

    for i, team in enumerate(data["teams"]):
        left = margin + (col_w + gap) * i
        header = f"{team['name']}（{team['tag']}）\n{team['outcome']}"
        body = _rows_to_text(team["rows"])
        _text(slide, left, top, col_w, Inches(0.42), header, size=6, bold=True, color=team["color"])
        _text(slide, left, top + Inches(0.42), col_w, Inches(1.55), body, size=5.5)

    cross = data["cross"]
    cross_top = Inches(2.6)
    _text(slide, margin, cross_top, w, Inches(0.35), f"{cross['name']}（{cross['tag']}）\n{cross['outcome']}", size=6, bold=True, color=cross["color"])

    half = (w - Inches(0.05)) / 2
    _text(slide, margin, cross_top + Inches(0.38), half, Inches(0.9), _rows_to_text(cross["rows_left"]), size=5.5)
    _text(slide, margin + half + Inches(0.05), cross_top + Inches(0.38), half, Inches(0.9), _rows_to_text(cross["rows_right"]), size=5.5)


def _strip_heavy_parts(src: Path, dst: Path) -> None:
    with zipfile.ZipFile(src, "r") as zin:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in SKIP_ZIP_ENTRIES:
                    continue
                if item.filename.startswith("ppt/printerSettings/"):
                    continue
                zout.writestr(item, zin.read(item.filename))
        dst.write_bytes(buf.getvalue())


def _mission_table(slide, left, top, width, height, rows_data: list[tuple[str, str, str]]):
    n_rows = 1 + len(rows_data)
    table = slide.shapes.add_table(n_rows, 3, left, top, width, height).table
    headers = ("ミッション", "現", "1年後")
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(6)
            p.font.bold = True
    for r, (mission, now, future) in enumerate(rows_data, start=1):
        for c, val in enumerate((mission, now, future)):
            cell = table.cell(r, c)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(5.5)
                p.font.bold = c == 0
    return table


def build_slide2_tables(slide) -> None:
    data = SLIDE2
    margin = Inches(0.2)
    w = Inches(10) - margin * 2

    _text(slide, margin, Inches(0.08), w, Inches(0.24), data["title"], size=14, bold=True)
    _text(slide, margin, Inches(0.3), w, Inches(0.14), data["subtitle"], size=5.5, color=RGBColor(0x55, 0x55, 0x55))

    col_w = (w - Inches(0.08)) / 3
    top = Inches(0.46)
    hdr_h = Inches(0.36)
    tbl_h = Inches(1.35)
    gap = Inches(0.04)

    for i, team in enumerate(data["teams"]):
        left = margin + (col_w + gap) * i
        header = f"{team['name']}（{team['tag']}）\n{team['outcome']}"
        _text(slide, left, top, col_w, hdr_h, header, size=5.5, bold=True, color=team["color"])
        _mission_table(slide, left, top + hdr_h, col_w, tbl_h, list(team["rows"]))

    cross = data["cross"]
    cross_top = top + hdr_h + tbl_h + Inches(0.06)
    _text(
        slide, margin, cross_top, w, Inches(0.28),
        f"{cross['name']}（{cross['tag']}）\n{cross['outcome']}",
        size=5.5, bold=True, color=cross["color"],
    )

    half = (w - Inches(0.04)) / 2
    tbl_top = cross_top + Inches(0.28)
    _mission_table(slide, margin, tbl_top, half, Inches(0.82), list(cross["rows_left"]))
    _mission_table(slide, margin + half + Inches(0.04), tbl_top, half, Inches(0.82), list(cross["rows_right"]))


def _save_prs(build_slide2_fn, out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    layout = prs.slide_layouts[6]
    build_slide1(prs.slides.add_slide(layout))
    build_slide2_fn(prs.slides.add_slide(layout))

    tmp = DIR / f"_{out_path.stem}-tmp.pptx"
    prs.save(tmp)
    _strip_heavy_parts(tmp, out_path)
    tmp.unlink(missing_ok=True)


def main() -> None:
    _save_prs(build_slide2, OUT_TEXT)
    _save_prs(build_slide2_tables, OUT_TABLES)

    print(f"Created: {OUT_TEXT.name}  （テキスト版・インポート向け）")
    print(f"Created: {OUT_TABLES.name}（表版・2枚目に表）")
    print()
    print("【表版のインポートが「アップロードできませんでした」になる場合】")
    print("  仕事用 Google では表入り PPTX のアップロードが弾かれることがあります。")
    print("  次のいずれかを試してください:")
    print()
    print("  1) マイドライブ経由（Upload タブを使わない）")
    print("     drive.google.com に product-dept-slides-tables.pptx を上げる")
    print("     → スライドで ファイル→スライドをインポート→「マイドライブ」タブから選択")
    print()
    print("  2) 既存スライドの2枚目に表を貼る（おすすめ）")
    print("     python3 build-mission-tables-tsv.py")
    print("     → mission-mapping-tables.tsv をスプレッドシートに貼る")
    print("     → 表の範囲をコピーしてスライド2枚目に貼り付け")
    print()
    print("  3) API でネイティブ作成: python3 create-google-slides-native.py")


if __name__ == "__main__":
    main()
